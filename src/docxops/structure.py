"""outline, add_heading, set_heading_numbering, add_paragraph_after, styles_apply."""

from __future__ import annotations

import re
from collections import Counter

from docx.oxml.ns import qn

from ._core import OUT, PATH, TEXT, Ctx, OpError, finish, load, obj, operation
from ._xml import (
    BUILTIN_STYLE_NAMES, PIC_NS, PPR_ORDER, STYLE_ORDER, V_NS, Styles, append_to_body,
    body_paragraphs, builtin_spec, find_heading, get_or_add_pPr, has_ancestor,
    headings, insert_ordered, is_w, iter_fields, make, paragraph, para_text,
    section_end_block, suppress_inherited_numbering, top_level_block,
)
from .toc import _inside_toc_sdt, find_existing_tocs

OCC = {"type": "integer", "minimum": 1}


def _mm(length) -> float | None:
    return None if length is None else round(length.mm, 1)


@operation(obj(
    "Describe document structure: headings tree, counts of tables/images/paragraphs, "
    "sections (page size, orientation, margins in mm), styles used, fields, TOC presence.",
    {"path": PATH},
))
def outline(ctx: Ctx, path: str) -> None:
    doc = load(path)
    body = doc.element.body
    styles = Styles(doc)

    tree: list[dict] = []
    stack: list[dict] = []
    flat = []
    for h in headings(doc, styles, skip=_inside_toc_sdt):
        node = {"level": h.level, "text": h.text, "style": h.style,
                "paragraph": h.index, "children": []}
        flat.append({k: v for k, v in node.items() if k != "children"})
        while stack and stack[-1]["level"] >= h.level:
            stack.pop()
        (stack[-1]["children"] if stack else tree).append(node)
        stack.append(node)

    paras = body_paragraphs(body)
    para_styles = Counter(styles.para_style_name(p) for p in paras)
    char_styles: Counter = Counter()
    for rs in body.iter(qn("w:rStyle")):
        char_styles[styles.name_of(rs.get(qn("w:val")))] += 1
    table_styles: Counter = Counter()
    for ts in body.iter(qn("w:tblStyle")):
        table_styles[styles.name_of(ts.get(qn("w:val")))] += 1

    images = sum(1 for _ in body.iter(f"{{{PIC_NS}}}pic")) + \
        sum(1 for _ in body.iter(f"{{{V_NS}}}imagedata"))
    tables = list(body.iter(qn("w:tbl")))
    fields = Counter(f.code for f in iter_fields(body) if f.code)

    sections = []
    for i, s in enumerate(doc.sections, 1):
        sections.append({
            "number": i,
            "orientation": "landscape" if s.orientation == 1 else "portrait",
            "start": str(s.start_type).split(".")[-1].split(" ")[0].lower(),
            "page_mm": [_mm(s.page_width), _mm(s.page_height)],
            "margins_mm": {"top": _mm(s.top_margin), "bottom": _mm(s.bottom_margin),
                           "left": _mm(s.left_margin), "right": _mm(s.right_margin)},
            "header_linked": s.header.is_linked_to_previous if i > 1 else False,
            "different_first_page": bool(s.different_first_page_header_footer),
        })

    ctx.details.update({
        "headings": tree,
        "headings_flat": flat,
        "counts": {
            "paragraphs": len(paras),
            "headings": len(flat),
            "tables": sum(1 for t in tables if not has_ancestor(t, "tbl")),
            "nested_tables": sum(1 for t in tables if has_ancestor(t, "tbl")),
            "images": images,
            "sections": len(sections),
            "comments": sum(1 for _ in body.iter(qn("w:commentRangeStart"))),
        },
        "sections": sections,
        "styles_used": {"paragraph": dict(para_styles.most_common()),
                        "character": dict(char_styles.most_common()),
                        "table": dict(table_styles.most_common())},
        "fields": dict(fields.most_common()),
        "has_toc": bool(find_existing_tocs(body, styles)),
    })


def _insert_after(anchor, elements: list) -> None:
    for el in elements:
        anchor.addnext(el)
        anchor = el


@operation(obj(
    "Add a heading. Default: end of document; with after_heading: at the end of that "
    "heading's section.",
    {"path": PATH, "text": TEXT, "level": {"type": "integer", "minimum": 1, "maximum": 9},
     "after_heading": TEXT, "occurrence": OCC, "out": OUT},
    ["path", "text"],
))
def add_heading(ctx: Ctx, path: str, text: str, level: int = 1,
                after_heading: str | None = None, occurrence: int = 1,
                out: str | None = None) -> None:
    doc = load(path)
    styles = Styles(doc)
    p = paragraph(text, styles.ensure(f"Heading {level}", ctx))
    if after_heading:
        h = find_heading(doc, after_heading, occurrence, ctx, styles)
        _insert_after(section_end_block(doc, h, styles), [p])
        ctx.details["placed"] = f"end of section {h.text!r}"
    else:
        append_to_body(doc.element.body, p)
        ctx.details["placed"] = "end of document"
    ctx.details["paragraph"] = body_paragraphs(doc.element.body).index(p)
    ctx.details["style"] = f"Heading {level}"
    ctx.changed = True
    finish(ctx, doc, path, out)


@operation(obj(
    "Insert paragraph(s) after a heading (where=heading) or at the end of its section "
    "(where=section_end). Newlines in text start new paragraphs.",
    {"path": PATH, "heading": TEXT, "text": {"type": "string"}, "style": TEXT,
     "where": {"enum": ["heading", "section_end"]}, "occurrence": OCC, "out": OUT},
    ["path", "heading", "text"],
))
def add_paragraph_after(ctx: Ctx, path: str, heading: str, text: str,
                        style: str | None = None, where: str = "heading",
                        occurrence: int = 1, out: str | None = None) -> None:
    doc = load(path)
    styles = Styles(doc)
    h = find_heading(doc, heading, occurrence, ctx, styles)
    sid = styles.ensure(style, ctx) if style else None
    lines = text.split("\n")
    new = [paragraph(line, sid) for line in lines]
    anchor = top_level_block(h.p, doc.element.body) if where == "heading" else \
        section_end_block(doc, h, styles)
    if where == "heading" and anchor is not h.p:
        anchor = h.p  # heading inside a table cell or content control
    _insert_after(anchor, new)
    ctx.details.update({"heading": h.text, "paragraphs_added": len(new), "where": where})
    ctx.changed = True
    finish(ctx, doc, path, out)


# ------------------------------------------------------- heading numbering

NUMBERING_NAME = "docxops-heading-numbering"
OUTLINE_FMT = [("upperRoman", "%1."), ("upperLetter", "%2."), ("decimal", "%3."),
               ("lowerLetter", "%4)"), ("lowerRoman", "%5)"), ("decimal", "(%6)"),
               ("lowerLetter", "(%7)"), ("lowerRoman", "(%8)"), ("decimal", "%9.")]


def _lvl_spec(scheme: str, i: int) -> tuple[str, str]:
    if scheme == "decimal":
        return "decimal", ".".join(f"%{k}" for k in range(1, i + 2))
    return OUTLINE_FMT[i]


def _sample_label(scheme: str, i: int) -> str:
    fmt, text = _lvl_spec(scheme, i)
    one = {"decimal": "1", "upperRoman": "I", "lowerRoman": "i",
           "upperLetter": "A", "lowerLetter": "a"}
    if scheme == "decimal":
        return ".".join("1" for _ in range(i + 1))
    return re.sub(r"%\d", one[fmt], text)


@operation(obj(
    "Number Heading 1..levels with a multilevel list linked to the heading styles "
    "(decimal: 1, 1.1, 1.1.1; outline: I., A., 1.). scheme=none removes it. Re-run replaces.",
    {"path": PATH, "scheme": {"enum": ["decimal", "outline", "none"]},
     "levels": {"type": "integer", "minimum": 1, "maximum": 9}, "out": OUT},
))
def set_heading_numbering(ctx: Ctx, path: str, scheme: str = "decimal", levels: int = 3,
                          out: str | None = None) -> None:
    doc = load(path)
    styles = Styles(doc)
    numbering = doc.part.numbering_part.element

    # Remove a previous run's list and unlink heading styles from it.
    old_abs = []
    for a in numbering.findall(qn("w:abstractNum")):
        name = a.find(qn("w:name"))
        if name is not None and name.get(qn("w:val")) == NUMBERING_NAME:
            old_abs.append(a.get(qn("w:abstractNumId")))
            numbering.remove(a)
    old_nums = set()
    for n in numbering.findall(qn("w:num")):
        ref = n.find(qn("w:abstractNumId"))
        if ref is not None and ref.get(qn("w:val")) in old_abs:
            old_nums.add(n.get(qn("w:numId")))
            numbering.remove(n)
    for lvl in range(1, 10):
        s = styles.find(f"Heading {lvl}")
        if s is None:
            continue
        numPr = s.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if numPr is None:
            continue
        nid = numPr.find(qn("w:numId"))
        if nid is not None and nid.get(qn("w:val")) in old_nums:
            numPr.getparent().remove(numPr)
        elif scheme != "none" and lvl <= levels:
            ctx.warn(f"Heading {lvl} already had list numbering; it was replaced")
            numPr.getparent().remove(numPr)
    ctx.details["removed_previous"] = bool(old_abs)
    if scheme == "none":
        ctx.changed = bool(old_abs)
        ctx.details["scheme"] = "none"
        finish(ctx, doc, path, out)
        return

    abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abs_id = max(abs_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    a = make("w:abstractNum", abstractNumId=abs_id)
    a.append(make("w:multiLevelType", val="multilevel"))
    a.append(make("w:name", val=NUMBERING_NAME))
    for i in range(9):
        fmt, text = _lvl_spec(scheme, i)
        lvl = make("w:lvl", ilvl=i)
        lvl.append(make("w:start", val=1))
        lvl.append(make("w:numFmt", val=fmt))
        if i < levels:
            lvl.append(make("w:pStyle", val=styles.ensure(f"Heading {i + 1}", ctx)))
        lvl.append(make("w:lvlText", val=text))
        lvl.append(make("w:lvlJc", val="left"))
        pPr = make("w:pPr")
        ind = 432 + 144 * i
        pPr.append(make("w:ind", left=ind, hanging=ind))
        lvl.append(pPr)
        a.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(a)
    else:
        cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
        if cleanup is not None:
            cleanup.addprevious(a)
        else:
            numbering.append(a)
    num = make("w:num", numId=num_id)
    num.append(make("w:abstractNumId", val=abs_id))
    last_num = numbering.findall(qn("w:num"))
    if last_num:
        last_num[-1].addnext(num)
    else:
        a.addnext(num)

    for i in range(levels):
        s = styles.find(f"Heading {i + 1}")
        pPr = s.find(qn("w:pPr"))
        if pPr is None:
            pPr = insert_ordered(s, make("w:pPr"), STYLE_ORDER)
        numPr = make("w:numPr")
        numPr.append(make("w:ilvl", val=i))
        numPr.append(make("w:numId", val=num_id))
        insert_ordered(pPr, numPr, PPR_ORDER)
    # TOC Heading is based on Heading 1; keep the TOC title unnumbered.
    suppress_inherited_numbering(styles, "TOCHeading")

    ctx.changed = True
    ctx.details.update({
        "scheme": scheme, "levels": levels, "numId": num_id, "abstractNumId": abs_id,
        "labels": [_sample_label(scheme, i) for i in range(levels)],
    })
    finish(ctx, doc, path, out)


# ------------------------------------------------------------ styles_apply

def _is_bold(run, para_bold: bool) -> bool:
    rPr = run.find(qn("w:rPr"))
    if rPr is not None:
        b = rPr.find(qn("w:b"))
        if b is not None:
            return b.get(qn("w:val")) not in ("0", "false", "off")
        rs = rPr.find(qn("w:rStyle"))
        if rs is not None and rs.get(qn("w:val")) == "Strong":
            return True
    return para_bold


def _style_bold(styles: Styles, sid: str | None) -> bool:
    s = styles.by_id.get(sid) if sid else styles.default_para
    seen = set()
    while s is not None and id(s) not in seen:
        seen.add(id(s))
        b = s.find(f"{qn('w:rPr')}/{qn('w:b')}")
        if b is not None:
            return b.get(qn("w:val")) not in ("0", "false", "off")
        based = s.find(qn("w:basedOn"))
        s = styles.by_id.get(based.get(qn("w:val"))) if based is not None else None
    return False


CLEARABLE = ("rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "color", "sz", "szCs", "u")

RULE = {
    "type": "object",
    "properties": {
        "style": TEXT, "regex": TEXT, "current_style": TEXT,
        "bold": {"type": "boolean"}, "max_chars": OCC,
    },
    "required": ["style"], "additionalProperties": False,
}


@operation(obj(
    "Set built-in styles on body paragraphs; first rule with all conditions true wins.",
    {"path": PATH, "rules": {"type": "array", "items": RULE},
     "clear_direct": {"type": "boolean"}, "dry_run": {"type": "boolean"}, "out": OUT},
    ["path", "rules"],
))
def styles_apply(ctx: Ctx, path: str, rules: list[dict], clear_direct: bool = False,
                 dry_run: bool = False, out: str | None = None) -> None:
    if not rules:
        raise OpError("rules is empty")
    doc = load(path)
    styles = Styles(doc)
    compiled = []
    for i, r in enumerate(rules):
        target = r["style"]
        existing = styles.find(target)
        if existing is not None and existing.get(qn("w:customStyle")) in ("1", "true"):
            raise OpError(f"rule {i + 1}: {target!r} is a custom style, not built-in")
        if existing is None and builtin_spec(target) is None:
            raise OpError(f"rule {i + 1}: unknown built-in style {target!r}. "
                          f"Known: {BUILTIN_STYLE_NAMES}")
        try:
            rx = re.compile(r["regex"]) if r.get("regex") else None
        except re.error as exc:
            raise OpError(f"rule {i + 1}: bad regex: {exc}") from None
        compiled.append((r, rx))

    changes = []
    counts: Counter = Counter()
    for idx, p in enumerate(body_paragraphs(doc.element.body)):
        if has_ancestor(p, "tbl") or _inside_toc_sdt(p):
            continue
        text = para_text(p).strip()
        if not text:
            continue
        sid = styles.para_style_id(p)
        cur = styles.name_of(sid)
        for n, (r, rx) in enumerate(compiled, 1):
            if rx is not None and not rx.search(text):
                continue
            if r.get("current_style") and cur.casefold() != r["current_style"].casefold():
                continue
            if "max_chars" in r and len(text) > r["max_chars"]:
                continue
            if "bold" in r:
                pb = _style_bold(styles, sid)
                runs = [x for x in p.iter(qn("w:r")) if (x.findtext(qn("w:t")) or "").strip()]
                if bool(runs) and all(_is_bold(x, pb) for x in runs) != r["bold"]:
                    continue
            target = r["style"]
            if cur.casefold() == target.casefold():
                break
            counts[target] += 1
            if len(changes) < 200:
                changes.append({"paragraph": idx, "text": text[:60], "from": cur,
                                "to": target, "rule": n})
            if not dry_run:
                new_id = styles.ensure(target, ctx)
                pPr = get_or_add_pPr(p)
                insert_ordered(pPr, make("w:pStyle", val=new_id), PPR_ORDER)
                if clear_direct:
                    for rPr in p.iter(qn("w:rPr")):
                        if is_w(rPr.getparent(), "r"):
                            for tag in CLEARABLE:
                                for el in rPr.findall(qn("w:" + tag)):
                                    rPr.remove(el)
            break
    total = sum(counts.values())
    ctx.details.update({"restyled": total, "by_style": dict(counts), "changes": changes,
                        "dry_run": dry_run})
    if total > len(changes):
        ctx.warn(f"changes list truncated to {len(changes)} of {total}")
    ctx.changed = bool(total) and not dry_run
    if not dry_run:
        finish(ctx, doc, path, out)
