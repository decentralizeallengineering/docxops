"""add_table and table_from_csv."""

from __future__ import annotations

import csv
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from ._core import OUT, PATH, TEXT, Ctx, OpError, finish, load, obj, operation
from ._xml import (
    Styles, append_to_body, find_heading, insert_ordered, is_w, make, section_end_block,
)

TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
              "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
              "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"]

HEX = {"type": "string", "pattern": "^[0-9A-Fa-f]{6}$"}


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _is_number(s: str) -> bool:
    try:
        float(s.replace(",", "").replace("%", "").strip())
        return bool(s.strip())
    except ValueError:
        return False


def build_table(ctx: Ctx, doc, rows: list[list], header: bool, style: str | None,
                header_fill: str | None, after_heading: str | None, occurrence: int) -> None:
    if not rows:
        raise OpError("rows is empty")
    width = max(len(r) for r in rows)
    if width == 0:
        raise OpError("rows have no cells")
    ragged = sum(1 for r in rows if len(r) != width)
    if ragged:
        ctx.warn(f"{ragged} row(s) were shorter than {width} cells and were padded")
    data = [[_cell_text(v) for v in r] + [""] * (width - len(r)) for r in rows]

    styles = Styles(doc)
    table = doc.add_table(rows=len(data), cols=width)
    tbl = table._tbl
    if style:
        s = styles.find(style, "table")
        if s is None:
            ctx.warn(f"table style {style!r} not found; table left unstyled")
        else:
            table.style = doc.styles[s.find(qn("w:name")).get(qn("w:val"))]
    look = tbl.tblPr.find(qn("w:tblLook"))
    if look is not None and not header:
        look.set(qn("w:firstRow"), "0")
        look.set(qn("w:val"), "0000")

    for r_i, values in enumerate(data):
        row = table.rows[r_i]
        for c_i, value in enumerate(values):
            cell = row.cells[c_i]
            par = cell.paragraphs[0]
            lines = value.split("\n")
            for k, line in enumerate(lines):
                rn = par.add_run(line)
                if k < len(lines) - 1:
                    rn.add_break()
                if header and r_i == 0:
                    rn.bold = True
            if not (header and r_i == 0) and _is_number(value):
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if header and r_i == 0 and header_fill:
                tcPr = cell._tc.get_or_add_tcPr()
                insert_ordered(tcPr, make("w:shd", val="clear", color="auto",
                                          fill=header_fill.upper()), TCPR_ORDER)
        if header and r_i == 0:
            trPr = row._tr.get_or_add_trPr()
            insert_ordered(trPr, make("w:tblHeader"), TRPR_ORDER)

    # doc.add_table appended at the end; move it if asked.
    body = doc.element.body
    if after_heading:
        h = find_heading(doc, after_heading, occurrence, ctx, styles)
        anchor = section_end_block(doc, h, styles)
        anchor.addnext(tbl)
        ctx.details["placed"] = f"end of section {h.text!r}"
    else:
        ctx.details["placed"] = "end of document"
    # Two adjacent tables merge in Word; keep a paragraph between them.
    prev = tbl.getprevious()
    if prev is not None and is_w(prev, "tbl"):
        tbl.addprevious(make("w:p"))
    nxt = tbl.getnext()
    if nxt is not None and is_w(nxt, "tbl"):
        tbl.addnext(make("w:p"))
    if tbl.getparent() is not body:  # pragma: no cover - defensive
        append_to_body(body, tbl)
    tables = [t for t in body.iter(qn("w:tbl"))]
    ctx.details.update({
        "rows": len(data), "columns": width, "header": header,
        "table_number": tables.index(tbl) + 1,
        "style": table.style.name if table.style is not None else None,
    })
    ctx.changed = True


@operation(obj(
    "Add a table (row 1 = bold repeating header unless header=false) at the end, or at "
    "the end of after_heading's section.",
    {"path": PATH, "rows": {"type": "array", "minItems": 1, "items": {
        "type": "array", "items": {"type": ["string", "number", "boolean", "null"]}}},
     "header": {"type": "boolean"}, "style": TEXT, "header_fill": HEX,
     "after_heading": TEXT, "occurrence": {"type": "integer", "minimum": 1}, "out": OUT},
    ["path", "rows"],
))
def add_table(ctx: Ctx, path: str, rows: list[list], header: bool = True,
              style: str = "Table Grid", header_fill: str | None = None,
              after_heading: str | None = None, occurrence: int = 1,
              out: str | None = None) -> None:
    doc = load(path)
    build_table(ctx, doc, rows, header, style, header_fill, after_heading, occurrence)
    finish(ctx, doc, path, out)


@operation(obj(
    "Add a table from a CSV file; options as in add_table.",
    {"path": PATH, "csv_path": PATH, "delimiter": {"type": "string", "maxLength": 1},
     "encoding": TEXT, "header": {"type": "boolean"}, "style": TEXT, "header_fill": HEX,
     "after_heading": TEXT, "max_rows": {"type": "integer", "minimum": 1}, "out": OUT},
    ["path", "csv_path"],
))
def table_from_csv(ctx: Ctx, path: str, csv_path: str, delimiter: str = ",",
                   encoding: str = "utf-8-sig", header: bool = True,
                   style: str = "Table Grid", header_fill: str | None = None,
                   after_heading: str | None = None, max_rows: int = 2000,
                   out: str | None = None) -> None:
    src = Path(csv_path)
    if not src.is_file():
        raise OpError(f"CSV not found: {csv_path}")
    try:
        with open(src, newline="", encoding=encoding) as fh:
            rows = [r for r in csv.reader(fh, delimiter=delimiter)]
    except (UnicodeDecodeError, LookupError) as exc:
        raise OpError(f"cannot read CSV with encoding {encoding!r}: {exc}") from None
    except csv.Error as exc:
        raise OpError(f"malformed CSV: {exc}") from None
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise OpError("CSV has no data")
    if len(rows) > max_rows:
        ctx.warn(f"CSV has {len(rows)} rows; only the first {max_rows} were used")
        rows = rows[:max_rows]
    doc = load(path)
    build_table(ctx, doc, rows, header, style, header_fill, after_heading, 1)
    ctx.details["source"] = str(src)
    finish(ctx, doc, path, out)
