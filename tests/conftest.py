import io
import struct
import zlib

import docx
import pytest
from docx.oxml.ns import qn


def png_bytes(w: int = 4, h: int = 3) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


@pytest.fixture
def make_doc(tmp_path):
    """Build a .docx from a list of (kind, value) items and return its path."""

    def build(items, name="doc.docx"):
        d = docx.Document()
        for kind, value in items:
            if kind == "title":
                d.add_paragraph(value, style="Title")
            elif kind.startswith("h"):
                d.add_heading(value, int(kind[1:]))
            elif kind == "p":
                d.add_paragraph(value)
            elif kind == "table":
                t = d.add_table(rows=len(value), cols=len(value[0]))
                for r, row in enumerate(value):
                    for c, v in enumerate(row):
                        t.cell(r, c).text = v
            elif kind == "image":
                d.add_picture(io.BytesIO(png_bytes()))
            elif kind == "runs":  # [(text, {"bold": True}), ...]
                p = d.add_paragraph()
                for text, fmt in value:
                    r = p.add_run(text)
                    for k, v in fmt.items():
                        setattr(r, k, v)
            else:
                raise ValueError(kind)
        path = tmp_path / name
        d.save(path)
        return str(path)

    return build


SAMPLE = [
    ("title", "Annual Report"),
    ("h1", "Introduction"),
    ("p", "Opening text."),
    ("h2", "Background"),
    ("p", "Some background."),
    ("h3", "Details"),
    ("h4", "Too deep for the TOC"),
    ("h1", "Results"),
    ("table", [["Crop", "Yield"], ["Kale", "12"]]),
    ("p", "See [[ref]] for numbers."),
    ("h2", "Discussion"),
    ("p", "Closing text."),
]


@pytest.fixture
def sample(make_doc):
    return make_doc(SAMPLE)


def instr_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:instrText")))


def field_instrs(el) -> list[str]:
    """Instruction strings of complex fields under el, in order."""
    out, stack = [], []
    for node in el.iter(qn("w:fldChar"), qn("w:instrText")):
        if node.tag == qn("w:instrText"):
            if stack:
                stack[-1] += node.text or ""
            continue
        kind = node.get(qn("w:fldCharType"))
        if kind == "begin":
            stack.append("")
        elif kind == "end" and stack:
            out.append(stack.pop().strip())
    return out
