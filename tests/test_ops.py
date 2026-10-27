import json

import docx
import pytest
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

import docxops
from conftest import SAMPLE, field_instrs


def ok(r):
    assert r["ok"], r
    json.dumps(r)
    assert set(r) == {"ok", "changed", "warnings", "details"}
    return r


def texts(path):
    return [p.text for p in docx.Document(path).paragraphs]


def styles_of(path):
    return [(p.style.name, p.text) for p in docx.Document(path).paragraphs]


# --------------------------------------------------------------------- outline

def test_outline(make_doc):
    path = make_doc(SAMPLE + [("image", None)])
    r = ok(docxops.outline(path))
    d = r["details"]
    assert r["changed"] is False
    tree = d["headings"]
    assert [n["text"] for n in tree] == ["Introduction", "Results"]
    assert tree[0]["children"][0]["text"] == "Background"
    assert tree[0]["children"][0]["children"][0]["children"][0]["text"] == "Too deep for the TOC"
    assert d["counts"]["tables"] == 1 and d["counts"]["images"] == 1
    assert d["counts"]["headings"] == 6 and d["counts"]["sections"] == 1
    assert d["sections"][0]["orientation"] == "portrait"
    assert d["styles_used"]["paragraph"]["Heading 1"] == 2
    assert d["has_toc"] is False
    ok(docxops.add_toc(path))
    d2 = docxops.outline(path)["details"]
    assert d2["has_toc"] is True and d2["fields"]["TOC"] == 1
    assert d2["counts"]["headings"] == 6  # TOC entries are not headings


# ----------------------------------------------------------------- add_heading

def test_add_heading(sample):
    ok(docxops.add_heading(sample, "Appendix", level=1))
    assert styles_of(sample)[-1] == ("Heading 1", "Appendix")
    r = ok(docxops.add_heading(sample, "Limitations", level=2, after_heading="Introduction"))
    t = texts(sample)
    # End of the Introduction section = just before "Results".
    assert t.index("Limitations") == t.index("Results") - 1
    assert r["details"]["paragraph"] == t.index("Limitations")
    assert docxops.add_heading(sample, "X", after_heading="Nope")["ok"] is False
    assert docxops.add_heading(sample, "X", level=12)["ok"] is False


# ------------------------------------------------------- set_heading_numbering

def _heading_numids(path):
    doc = docx.Document(path)
    out = {}
    for i in range(1, 4):
        s = doc.styles[f"Heading {i}"].element
        num = s.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        out[i] = num.get(qn("w:val")) if num is not None else None
    return doc, out


def test_set_heading_numbering(sample):
    r = ok(docxops.set_heading_numbering(sample, scheme="decimal", levels=3))
    assert r["details"]["labels"] == ["1", "1.1", "1.1.1"]
    doc, ids = _heading_numids(sample)
    assert len(set(ids.values())) == 1 and ids[1] is not None
    numbering = doc.part.numbering_part.element
    ours = [a for a in numbering.findall(qn("w:abstractNum"))
            if a.find(qn("w:name")) is not None
            and a.find(qn("w:name")).get(qn("w:val")) == "docxops-heading-numbering"]
    assert len(ours) == 1
    lvl2 = ours[0].findall(qn("w:lvl"))[1]
    assert lvl2.find(qn("w:lvlText")).get(qn("w:val")) == "%1.%2"
    assert lvl2.find(qn("w:pStyle")).get(qn("w:val")) == "Heading2"

    # Re-run replaces rather than stacking a second list.
    ok(docxops.set_heading_numbering(sample, scheme="outline", levels=2))
    doc, ids = _heading_numids(sample)
    numbering = doc.part.numbering_part.element
    names = [a.find(qn("w:name")) for a in numbering.findall(qn("w:abstractNum"))]
    assert sum(1 for n in names if n is not None
               and n.get(qn("w:val")) == "docxops-heading-numbering") == 1
    assert ids[3] is None and ids[1] == ids[2]

    ok(docxops.set_heading_numbering(sample, scheme="none"))
    _, ids = _heading_numids(sample)
    assert set(ids.values()) == {None}


# --------------------------------------------------------- add_paragraph_after

def test_add_paragraph_after(sample):
    ok(docxops.add_paragraph_after(sample, "Background", "Right after.\nSecond line."))
    t = texts(sample)
    i = t.index("Background")
    assert t[i + 1:i + 3] == ["Right after.", "Second line."]
    # A section includes its subsections: Background's ends just before "Results".
    ok(docxops.add_paragraph_after(sample, "background", "End of background.",
                                   where="section_end", style="Quote"))
    t = texts(sample)
    assert t.index("End of background.") == t.index("Results") - 1
    assert t.index("End of background.") > t.index("Too deep for the TOC")
    assert ("Quote", "End of background.") in styles_of(sample)
    bad = docxops.add_paragraph_after(sample, "Nowhere", "x")
    assert bad["ok"] is False and "Headings" in bad["details"]["error"]


# ---------------------------------------------------------------- replace_text

def test_replace_text_across_runs_keeps_formatting(make_doc):
    path = make_doc([
        ("runs", [("Hel", {"bold": True}), ("lo Wor", {"italic": True}),
                  ("ld", {}), ("!", {"underline": True})]),
        ("runs", [("keep ", {"bold": True}), ("me", {"italic": True})]),
        ("p", "alpha\tbeta"),
    ])
    d = docx.Document(path)
    d.paragraphs[2].runs[0].text = "alpha"
    d.paragraphs[2].runs[0].add_tab()
    d.paragraphs[2].runs[0].add_text("beta")
    d.save(path)

    r = ok(docxops.replace_text(path, "Hello World", "Hi Earth"))
    assert r["details"]["replaced"] == 1 and r["details"]["mixed_formatting"] == 1
    doc = docx.Document(path)
    p0 = doc.paragraphs[0]
    assert p0.text == "Hi Earth!"
    assert p0.runs[0].bold and p0.runs[0].text == "Hi Earth"
    assert p0.runs[-1].text == "!" and p0.runs[-1].underline  # untouched run unchanged
    p1 = doc.paragraphs[1]
    assert [(x.text, x.bold, x.italic) for x in p1.runs] == [("keep ", True, None), ("me", None, True)]

    # A match across a tab is reported, not forced.
    assert ok(docxops.replace_text(path, "alpha beta", "x"))["details"]["replaced"] == 0
    r = ok(docxops.replace_text(path, "a\tb", "x"))
    assert r["details"]["replaced"] == 0 and r["details"]["skipped_count"] == 1
    assert r["changed"] is False and r["warnings"]

    r = ok(docxops.replace_text(path, "HI", "Hey", match_case=False, whole_word=True))
    assert r["details"]["replaced"] == 1
    assert docx.Document(path).paragraphs[0].text == "Hey Earth!"


def test_replace_text_reports_match_across_field_result(make_doc):
    path = make_doc([("table", [["a"]])])
    assert docxops.add_caption(path, "Table", "Yields")["ok"]
    r = ok(docxops.replace_text(path, "Table 1", "Tbl 9"))
    assert r["details"]["replaced"] == 0 and r["details"]["skipped_count"] == 1
    assert "field" in r["details"]["skipped"][0]["reason"]
    assert ok(docxops.replace_text(path, "Yields", "Harvest"))["details"]["replaced"] == 1
    cap = next(p for p in docx.Document(path).paragraphs if p.style.name == "Caption")
    assert cap.text == "Table 1: Harvest" and field_instrs(cap._p) == ["SEQ Table \\* ARABIC"]


def test_replace_text_headers_and_limit(make_doc):
    path = make_doc([("p", "cat cat cat")])
    d = docx.Document(path)
    d.sections[0].header.paragraphs[0].text = "cat header"
    d.save(path)
    r = ok(docxops.replace_text(path, "cat", "dog", max_replacements=2))
    assert r["details"]["replaced"] == 2
    assert texts(path)[0] == "dog dog cat"
    ok(docxops.replace_text(path, "cat", "dog", scope="all"))
    d = docx.Document(path)
    assert d.sections[0].header.paragraphs[0].text == "dog header"


# ------------------------------------------------------------------- add_table

def test_add_table(sample):
    rows = [["Crop", "Beds", "Notes"], ["Kale", 4, None], ["Chard", 2.0, True]]
    r = ok(docxops.add_table(sample, rows, after_heading="Introduction", header_fill="D9E2F3"))
    assert r["details"]["rows"] == 3 and r["details"]["columns"] == 3
    doc = docx.Document(sample)
    t = doc.tables[0]  # placed before the existing Results table
    assert [c.text for c in t.rows[1].cells] == ["Kale", "4", ""]
    assert [c.text for c in t.rows[2].cells] == ["Chard", "2", "TRUE"]
    assert t.style.name == "Table Grid"
    assert all(run.bold for c in t.rows[0].cells for run in c.paragraphs[0].runs)
    assert t.rows[0]._tr.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None
    shd = t.rows[0].cells[0]._tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}")
    assert shd.get(qn("w:fill")) == "D9E2F3"
    body = list(doc.element.body)
    assert body[body.index(t._tbl) + 1].tag == qn("w:p")  # before "Results" heading

    r = ok(docxops.add_table(sample, [["a", "b"], ["c"]], header=False, style="No Such"))
    assert len(r["warnings"]) == 2
    assert docxops.add_table(sample, [])["ok"] is False


def test_table_from_csv(sample, tmp_path):
    csv_path = tmp_path / "harvest.csv"
    csv_path.write_text("Crop;Lbs\nKale;12\nBeets;30\n", encoding="utf-8")
    r = ok(docxops.table_from_csv(sample, str(csv_path), delimiter=";", max_rows=2))
    assert r["details"]["rows"] == 2 and r["warnings"]
    t = docx.Document(sample).tables[-1]
    assert [[c.text for c in row.cells] for row in t.rows] == [["Crop", "Lbs"], ["Kale", "12"]]
    assert docxops.table_from_csv(sample, str(tmp_path / "none.csv"))["ok"] is False


# ----------------------------------------------------------------- add_caption

def _captions(path):
    doc = docx.Document(path)
    return [(p.text, field_instrs(p._p)) for p in doc.paragraphs if p.style.name == "Caption"]


def test_add_caption_tables_and_figures(make_doc):
    path = make_doc([("table", [["a"]]), ("p", "between"), ("table", [["b"]]),
                     ("image", None)])
    r = ok(docxops.add_caption(path, "Table", "Second table", target=2))
    assert r["details"]["label"] == "Table 1"
    r = ok(docxops.add_caption(path, "Table", "First table", target=1))
    assert r["details"]["label"] == "Table 1" and r["details"]["bookmark"].startswith("_Ref")
    caps = _captions(path)
    assert [c[0] for c in caps] == ["Table 1: First table", "Table 2: Second table"]
    assert all(c[1] == ["SEQ Table \\* ARABIC"] for c in caps)

    doc = docx.Document(path)
    body = list(doc.element.body)
    assert body[body.index(doc.tables[0]._tbl) - 1].tag == qn("w:p")
    assert doc.tables[0]._tbl.getprevious().find(f"{qn('w:pPr')}/{qn('w:pStyle')}") \
        .get(qn("w:val")) == "Caption"

    # Running again on the same table replaces its caption.
    r = ok(docxops.add_caption(path, "Table", "First table, revised", target=1))
    assert r["details"]["replaced"] is True
    assert [c[0] for c in _captions(path)][:2] == ["Table 1: First table, revised",
                                                    "Table 2: Second table"]

    r = ok(docxops.add_caption(path, "Figure", "Red pixels"))
    assert r["details"]["label"] == "Figure 1" and r["details"]["position"] == "below"
    assert docxops.add_caption(path, "Table", target=9)["ok"] is False


# ------------------------------------------------------------- cross_reference

def test_cross_reference_caption_and_heading(sample):
    ok(docxops.add_caption(sample, "Table", "Yields", target=1))
    r = ok(docxops.cross_reference(sample, "Table 1", "[[ref]]"))
    assert r["details"]["target_kind"] == "caption" and r["details"]["result"] == "Table 1"
    doc = docx.Document(sample)
    para = next(p for p in doc.paragraphs if p.text.startswith("See "))
    assert para.text == "See Table 1 for numbers."
    bm = r["details"]["bookmark"]
    assert f"REF {bm} \\h" in field_instrs(para._p)
    cap = next(p for p in doc.paragraphs if p.style.name == "Caption")
    assert bm in [b.get(qn("w:name")) for b in cap._p.iter(qn("w:bookmarkStart"))]

    ok(docxops.add_paragraph_after(sample, "Discussion", "Method is on page {pg}, {where}."))
    r = ok(docxops.cross_reference(sample, "Background", "{pg}", display="page"))
    assert r["details"]["target_kind"] == "heading" and r["details"]["result"] == ""
    r = ok(docxops.cross_reference(sample, "Background", "{where}", display="above_below"))
    assert r["details"]["result"] == "above"
    doc = docx.Document(sample)
    para = next(p for p in doc.paragraphs if p.text.startswith("Method"))
    assert para.text == "Method is on page , above."
    assert any(i.startswith("PAGEREF _Ref") for i in field_instrs(para._p))
    assert doc.settings.element.find(qn("w:updateFields")) is not None

    assert docxops.cross_reference(sample, "Table 7", "x", target_type="caption")["ok"] is False
    assert docxops.cross_reference(sample, "Results", "[[missing]]")["ok"] is False


# --------------------------------------------------------------- header_footer

def test_header_footer(sample):
    r = ok(docxops.header_footer(sample, header="Jana Farms - {page}", page_numbers="page_x_of_y",
                                 align="right"))
    assert {"section": 1, "part": "footer"} in r["details"]["updated"]
    doc = docx.Document(sample)
    footer = doc.sections[0].footer.paragraphs[0]
    assert footer.text == "Page  of "
    assert [i.split()[0] for i in field_instrs(footer._p)] == ["PAGE", "NUMPAGES"]
    header = doc.sections[0].header.paragraphs[0]
    assert header.text == "Jana Farms - " and field_instrs(header._p) == ["PAGE \\* MERGEFORMAT"]
    assert header.alignment == 2  # right
    # Re-running replaces content instead of appending.
    ok(docxops.header_footer(sample, footer="Confidential"))
    assert [p.text for p in docx.Document(sample).sections[0].footer.paragraphs] == ["Confidential"]
    assert docxops.header_footer(sample)["ok"] is False


# --------------------------------------------------------------- section_break

def test_section_break_orientation(sample):
    r = ok(docxops.section_break(sample, before_heading="Background", orientation="landscape",
                                 isolate=True))
    assert r["details"]["section"] == 2 and r["details"]["sections"] == 3
    doc = docx.Document(sample)
    orients = [s.orientation for s in doc.sections]
    assert orients == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
    s2 = doc.sections[1]
    assert s2.page_width > s2.page_height
    # Background starts section 2 and Results starts section 3: the paragraph
    # before each carries the sectPr of the section it ends.
    for heading in ("Background", "Results"):
        p = next(x for x in doc.paragraphs if x.text == heading)
        assert p._p.getprevious().find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    # Section 3 runs to the end of the document; a heading whose section does too
    # cannot be isolated, which is reported.
    r = ok(docxops.section_break(sample, before_heading="Results", isolate=True))
    assert r["warnings"]

    r = ok(docxops.section_break(sample, start="continuous"))
    assert docx.Document(sample).sections[-1].start_type == 0  # continuous
    assert docxops.section_break(sample, before_heading="Annual Report")["ok"] is False


# ------------------------------------------------------------------ page_setup

def test_page_setup(sample):
    r = ok(docxops.page_setup(sample, paper="A4", orientation="landscape",
                              margins={"top": 20, "left": 25.4, "right": 25.4, "bottom": 20}))
    assert r["details"]["sections"][0]["page_mm"] == [297.0, 210.0]
    s = docx.Document(sample).sections[0]
    assert s.orientation == WD_ORIENT.LANDSCAPE
    assert s.left_margin.twips == 1440 and abs(s.top_margin.mm - 20) < 0.05
    ok(docxops.page_setup(sample, orientation="portrait"))
    s = docx.Document(sample).sections[0]
    assert round(s.page_width.mm) == 210
    assert docxops.page_setup(sample, margins={"left": 200, "right": 200})["ok"] is False
    assert docxops.page_setup(sample, margins={"middle": 1})["ok"] is False
    assert docxops.page_setup(sample, section=5, paper="A4")["ok"] is False


# --------------------------------------------------------------------- compose

def test_compose_keeps_styles_and_numbering(make_doc, tmp_path):
    a = make_doc([("h1", "Part A"), ("p", "alpha")], "a.docx")
    b = make_doc([("h1", "Part B")], "b.docx")
    d = docx.Document(b)
    from docx.enum.style import WD_STYLE_TYPE
    st = d.styles.add_style("Field Note", WD_STYLE_TYPE.PARAGRAPH)
    st.font.bold = True
    d.add_paragraph("note", style="Field Note")
    d.add_paragraph("first item", style="List Number")
    d.add_paragraph("second item", style="List Number")
    d.save(b)

    out = str(tmp_path / "merged.docx")
    r = ok(docxops.compose(a, [b], out=out))
    assert r["details"]["appended"][0]["path"] == b
    merged = docx.Document(out)
    t = [p.text for p in merged.paragraphs]
    assert t.index("Part A") < t.index("Part B")
    assert ("Field Note", "note") in [(p.style.name, p.text) for p in merged.paragraphs]
    assert "List Number" in [p.style.name for p in merged.paragraphs]
    assert any(br.get(qn("w:type")) == "page" for br in merged.element.body.iter(qn("w:br")))
    assert docx.Document(a).paragraphs[-1].text == "alpha"  # source untouched
    assert docxops.compose(a, [str(tmp_path / "nope.docx")])["ok"] is False


# ---------------------------------------------------------------- styles_apply

def test_styles_apply(make_doc):
    path = make_doc([
        ("runs", [("FIELD NOTES", {"bold": True})]),
        ("p", "Ordinary body text that is quite long and not a heading at all."),
        ("runs", [("Chapter 2: Soil", {"bold": True})]),
        ("p", "Note: water at dawn"),
    ])
    rules = [
        {"style": "Heading 1", "bold": True, "max_chars": 40},
        {"style": "Quote", "regex": "^Note:"},
    ]
    r = ok(docxops.styles_apply(path, rules, dry_run=True))
    assert r["details"]["restyled"] == 3 and r["changed"] is False
    assert styles_of(path)[0][0] == "Normal"

    r = ok(docxops.styles_apply(path, rules, clear_direct=True))
    assert r["details"]["by_style"] == {"Heading 1": 2, "Quote": 1}
    got = styles_of(path)
    assert [g[0] for g in got] == ["Heading 1", "Normal", "Heading 1", "Quote"]
    assert docx.Document(path).paragraphs[0].runs[0].bold is None  # direct bold cleared
    assert docxops.styles_apply(path, [{"style": "Fancy Custom"}])["ok"] is False
    assert docxops.styles_apply(path, [{"style": "Title", "regex": "("}])["ok"] is False


# --------------------------------------------------------------- comments_list

def test_comments_list(make_doc):
    path = make_doc([("p", "Check the irrigation timer."), ("p", "No comment here.")])
    d = docx.Document(path)
    run = d.paragraphs[0].runs[0]
    d.add_comment(run, text="Timer is set to 6am.", author="Olive", initials="OS")
    d.save(path)
    r = ok(docxops.comments_list(path))
    assert r["details"]["count"] == 1
    c = r["details"]["comments"][0]
    assert c["author"] == "Olive" and c["initials"] == "OS"
    assert c["text"] == "Timer is set to 6am."
    assert c["anchor_text"] == "Check the irrigation timer."
    assert ok(docxops.comments_list(make_doc([("p", "x")], "plain.docx")))["details"]["count"] == 0


# ---------------------------------------------------------- contract / schemas

def test_schemas_and_registry():
    import jsonschema
    assert set(docxops.SCHEMAS) == set(docxops.OPERATIONS)
    assert set(docxops.__all__) >= set(docxops.OPERATIONS)
    assert len(docxops.OPERATIONS) == 16
    for name, schema in docxops.SCHEMAS.items():
        assert len(json.dumps(schema, separators=(",", ":"))) <= 600, name
        jsonschema.Draft202012Validator.check_schema(schema)
        assert schema["type"] == "object" and "path" in schema["required"]
    tools = docxops.tool_definitions()
    assert {t["name"] for t in tools} == set(docxops.SCHEMAS)


@pytest.mark.parametrize("name", sorted(docxops.SCHEMAS))
def test_every_operation_reports_instead_of_raising(name, tmp_path):
    fn = docxops.OPERATIONS[name]
    r = fn(path=str(tmp_path / "missing.docx"))  # missing file and/or missing args
    assert r["ok"] is False and r["changed"] is False and r["details"]["error"]
    json.dumps(r)
    garbage = tmp_path / "garbage.docx"
    garbage.write_bytes(b"not a zip")
    r = docxops.call(name, {"path": str(garbage), "bogus_arg": 1})
    assert r["ok"] is False


def test_call_dispatch(sample):
    r = docxops.call("outline", {"path": sample})
    assert r["ok"] and r["details"]["counts"]["tables"] == 1
    assert docxops.call("nope", {})["ok"] is False
    assert docxops.call("add_heading", {"path": sample, "text": "X", "level": "one"})["ok"] is False
