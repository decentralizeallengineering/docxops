import json

import docx
from docx.oxml.ns import qn

import docxops
from conftest import field_instrs

TOC_INSTR = 'TOC \\o "1-3" \\h \\z \\u'


def toc_sdts(doc):
    out = []
    for sdt in doc.element.body.iter(qn("w:sdt")):
        gal = sdt.find(f"{qn('w:sdtPr')}/{qn('w:docPartObj')}/{qn('w:docPartGallery')}")
        if gal is not None and gal.get(qn("w:val")) == "Table of Contents":
            out.append(sdt)
    return out


def toc_entries(sdt):
    """(style id, text, anchor) for each hyperlinked entry paragraph."""
    entries = []
    for p in sdt.iter(qn("w:p")):
        link = p.find(qn("w:hyperlink"))
        if link is None:
            continue
        style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val"))
        text = "".join(t.text for t in link.iter(qn("w:t")))
        entries.append((style, text, link.get(qn("w:anchor"))))
    return entries


def heading_bookmarks(doc):
    """{heading text: [bookmark names starting with _Toc]}"""
    out = {}
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            names = [b.get(qn("w:name")) for b in p._p.findall(qn("w:bookmarkStart"))]
            out[p.text] = [n for n in names if n.startswith("_Toc")]
    return out


def test_add_toc_field_entries_bookmarks_and_rerun(sample):
    r = docxops.add_toc(sample)
    assert r["ok"] and r["changed"], r
    json.dumps(r)
    assert r["details"]["field_code"] == TOC_INSTR

    doc = docx.Document(sample)
    sdts = toc_sdts(doc)
    assert len(sdts) == 1
    sdt = sdts[0]

    # Real TOC field: begin/instr/separate ... end, with the expected code.
    instrs = field_instrs(sdt)
    assert TOC_INSTR in instrs
    kinds = [f.get(qn("w:fldCharType")) for f in sdt.iter(qn("w:fldChar"))]
    assert kinds[0] == "begin" and kinds[-1] == "end"
    assert kinds.count("begin") == kinds.count("end")

    # Pre-filled result: one hyperlinked entry per heading in levels 1-3, TOC n styles.
    entries = toc_entries(sdt)
    assert [(s, t) for s, t, _ in entries] == [
        ("TOC1", "Introduction"), ("TOC2", "Background"), ("TOC3", "Details"),
        ("TOC1", "Results"), ("TOC2", "Discussion"),
    ]
    # Word stores built-in TOC style names in lower case ("toc 1").
    style_names = {s.style_id: s.name for s in doc.styles}
    assert style_names["TOC1"] == "toc 1" and style_names["TOC3"] == "toc 3"

    # Title paragraph uses TOC Heading (outline level 9, so it never lists itself).
    first_p = sdt.find(qn("w:sdtContent")).find(qn("w:p"))
    assert first_p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val")) == "TOCHeading"
    assert "".join(t.text for t in first_p.iter(qn("w:t"))) == "Contents"

    # Each entry points at a bookmark that sits on its heading.
    bms = heading_bookmarks(doc)
    for _, text, anchor in entries:
        assert bms[text] == [anchor]
    assert bms["Too deep for the TOC"] == []

    # Page numbers are PAGEREF fields with an empty result, never a guessed number.
    pagerefs = [i for i in instrs if i.startswith("PAGEREF")]
    assert len(pagerefs) == len(entries)
    for link in sdt.iter(qn("w:hyperlink")):
        texts = [t.text for t in link.iter(qn("w:t"))]
        assert len(texts) == 1  # only the heading text, no page number

    # TOC sits after the Title paragraph; Word refreshes fields on open.
    body = list(doc.element.body)
    assert body[1] is sdt
    uf = doc.settings.element.find(qn("w:updateFields"))
    assert uf is not None and uf.get(qn("w:val")) == "true"

    # Second run: replaced in place, not duplicated; bookmarks reused, not stacked.
    r2 = docxops.add_toc(sample, levels="1-2", title="Table of Contents")
    assert r2["ok"], r2
    assert r2["details"]["replaced_existing"] == 1
    assert r2["details"]["bookmarks_added"] == 0
    doc2 = docx.Document(sample)
    sdts2 = toc_sdts(doc2)
    assert len(sdts2) == 1
    assert sum(1 for i in field_instrs(doc2.element.body) if i.startswith("TOC")) == 1
    assert 'TOC \\o "1-2" \\h \\z \\u' in field_instrs(sdts2[0])
    assert [t for _, t, _ in toc_entries(sdts2[0])] == [
        "Introduction", "Background", "Results", "Discussion"]
    assert list(doc2.element.body)[1] is sdts2[0]
    for names in heading_bookmarks(doc2).values():
        assert len(names) <= 1
    assert len(doc2.settings.element.findall(qn("w:updateFields"))) == 1

    # Third identical run changes nothing structurally.
    docxops.add_toc(sample, levels="1-2", title="Table of Contents")
    doc3 = docx.Document(sample)
    assert len(toc_sdts(doc3)) == 1
    assert toc_entries(toc_sdts(doc3)[0]) == toc_entries(sdts2[0])


def test_add_toc_replaces_word_style_toc_without_content_control(make_doc, tmp_path):
    path = make_doc([("p", "placeholder"), ("h1", "Alpha"), ("h2", "Beta")])
    d = docx.Document(path)
    p = d.paragraphs[0]
    p.text = ""
    from docxops._xml import field_runs
    for r in field_runs('TOC \\o "1-3" \\h \\z \\u', "Old stale entry"):
        p._p.append(r)
    d.save(path)

    r = docxops.add_toc(path)
    assert r["ok"] and r["details"]["replaced_existing"] == 1, r
    doc = docx.Document(path)
    assert "Old stale entry" not in "".join(t.text for t in doc.element.body.iter(qn("w:t")))
    assert sum(1 for i in field_instrs(doc.element.body) if i.startswith("TOC")) == 1
    assert [t for _, t, _ in toc_entries(toc_sdts(doc)[0])] == ["Alpha", "Beta"]


def test_add_toc_includes_direct_outline_level_and_leaves_table_of_figures(make_doc):
    path = make_doc([("h1", "One"), ("p", "Outlined body paragraph")])
    d = docx.Document(path)
    from docxops._xml import field_runs
    pPr = d.paragraphs[1]._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), "1")
    pPr.append(ol)
    tof = d.add_paragraph()
    for r in field_runs('TOC \\h \\z \\c "Figure"', ""):
        tof._p.append(r)
    d.save(path)

    r = docxops.add_toc(path, levels=2, position="start")
    assert r["ok"], r
    doc = docx.Document(path)
    assert list(doc.element.body)[0].tag == qn("w:sdt")
    entries = toc_entries(toc_sdts(doc)[0])
    assert [(s, t) for s, t, _ in entries] == [("TOC1", "One"), ("TOC2", "Outlined body paragraph")]
    # The table of figures is a different TOC and must survive.
    assert any("\\c" in i for i in field_instrs(doc.element.body))


def test_add_toc_errors_and_empty_documents(make_doc, tmp_path):
    bad = docxops.add_toc(str(tmp_path / "missing.docx"))
    assert bad["ok"] is False and "not found" in bad["details"]["error"]
    assert docxops.add_toc(make_doc([("h1", "A")]), levels="3-1")["ok"] is False
    assert docxops.add_toc(make_doc([("h1", "A")], "b.docx"), levels="x")["ok"] is False

    empty = make_doc([("p", "no headings here")], "empty.docx")
    r = docxops.add_toc(empty)
    assert r["ok"] and r["warnings"]
    assert TOC_INSTR in field_instrs(docx.Document(empty).element.body)


def test_add_toc_out_leaves_source_untouched(sample, tmp_path):
    out = str(tmp_path / "copy.docx")
    r = docxops.add_toc(sample, out=out)
    assert r["ok"] and r["details"]["saved_to"] == out
    assert not toc_sdts(docx.Document(sample))
    assert toc_sdts(docx.Document(out))


def test_toc_title_stays_unnumbered_when_headings_are_numbered(sample):
    # TOC Heading is based on Heading 1; numbering Heading 1 must not number the TOC title.
    assert docxops.set_heading_numbering(sample)["ok"]
    assert docxops.add_toc(sample)["ok"]
    doc = docx.Document(sample)
    toc_heading = doc.styles["TOC Heading"].element
    num = toc_heading.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
    assert num is not None and num.get(qn("w:val")) == "0"
    h1 = doc.styles["Heading 1"].element.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
    assert h1 is not None and h1.get(qn("w:val")) != "0"
